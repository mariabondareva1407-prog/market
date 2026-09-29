#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты правок 09.2026 по шаблонам: общие B2B-колонки, единый реестр заглушек
с отчётом и подтверждением, LockedInfo в пространстве id ПСБ, общий
разделитель изображений, штрихкод ровно 13 символов.

Сети нет: apigw и категорийная API Белки подменяются заглушками, данные —
реальные (файл ассортимента Мвидео и JSON-экспорт Базара из архива).

Запуск:  python3 tests/test_placeholders_and_columns.py
"""
from __future__ import annotations

import glob
import pathlib
import io
import json
import os
import shutil
import sys
import tempfile
import traceback
from contextlib import redirect_stdout
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
PROJECT = Path(__file__).resolve().parent.parent

from openpyxl import load_workbook                                  # noqa: E402
from belka.context import PipelineContext, WarehouseConfig          # noqa: E402
from belka.core.columns import build_barcode, normalize_attr_key, IMAGES_SEPARATOR  # noqa: E402
from belka.core import defaults                                     # noqa: E402
from belka.stages import psb_generate as psb_mod                    # noqa: E402
from belka.warehouses import bazar, mvideo                          # noqa: E402

RESULTS = []
B2B = ["ТН ВЭД *", "Кратность заказа *", "Минимальное количество *", "Количество в упаковке *"]

FAKE_CATEGORY = {"properties": [
    {"display_name": "Бренд(Brend)", "is_required": True, "has_directory": True},
    {"display_name": "Тип товара", "is_required": False, "has_directory": True},
]}


def check(name, fn):
    try:
        fn()
    except Exception:
        RESULTS.append((name, False)); print(f"✘ {name}"); traceback.print_exc()
    else:
        RESULTS.append((name, True)); print(f"✔ {name}")


# ------------------------------------------------------------ базовые --

def test_barcode_exactly_13():
    for raw in ("49440", "400457846", "20086324", 1140):
        bc = build_barcode("4700", raw)
        assert len(bc) == 13, (raw, bc, len(bc))
        assert bc.startswith("4700")
    # длинный id — берём последние цифры, соседние остаются различимыми
    a, b = build_barcode("4700", "1234567890123"), build_barcode("4700", "1234567890124")
    assert len(a) == len(b) == 13 and a != b
    assert build_barcode("4700", "") == ""


def test_normalize_attr_key():
    for variant in ("Бренд(Brend) * (справочник)", "Бренд(Brend)", "Бренд", "БРЕНД (справочник)"):
        assert normalize_attr_key(variant) == "бренд", variant
    assert normalize_attr_key("Цвет * (справочник)") == "цвет"


def test_placeholder_registry_and_override():
    ctx = PipelineContext(config=bazar.CONFIG, env={}, interactive=False, assume_yes=True)
    assert ctx.placeholder("tnved") == "8500000000"
    assert ctx.placeholder("description") == "нет описания"
    assert ctx.placeholder("activity") == "да"
    for key in ("order_multiplicity", "min_order_quantity", "quantity_in_package"):
        assert ctx.placeholder(key) == 1, key

    cfg = WarehouseConfig(key="t", display_name="t", seller_id=6, store_id=1,
                          placeholders={"description": "Нет данных", "tnved": None})
    ctx2 = PipelineContext(config=cfg, env={}, interactive=False, assume_yes=True)
    assert ctx2.placeholder("description") == "Нет данных"
    assert ctx2.placeholder("tnved") is None          # «этому складу не подставлять»
    assert ctx2.placeholder("activity") == "да"        # остальное наследуется
    assert any("рецепт склада" in l for l in ctx2.placeholders_report())


def test_every_placeholder_column_is_known():
    """Колонка в реестре должна писаться так же, как в шаблоне (после clean_header)."""
    known = set(B2B) | {"Описание *", "НДС *", "Активность товара на маркете *", "Штрихкод *",
                        "Вес, Кг *", "Ширина, см *", "Высота, см *", "Длина, см *"}
    known_clean = {h.replace(" *", "") for h in known}
    for ph in defaults.PLACEHOLDERS.values():
        for col in ph.columns:
            assert col in known_clean, f"{col!r} из реестра не совпадает ни с одной колонкой шаблона"


# ------------------------------------------------- Мвидео, на реальных данных --

def _fake_psb_env(tmp: Path):
    """Готовит рабочую папку и подменяет всё сетевое. Возвращает ctx."""
    for name in ("console_to_admin_category_map.json",):
        shutil.copy(PROJECT / name, tmp / name)
    bridge = json.load(open(tmp / "console_to_admin_category_map.json", encoding="utf-8"))
    psb_ids = [int(k) for k in list(bridge)[:3]]

    ctx = PipelineContext(config=mvideo.CONFIG, env={}, interactive=False, assume_yes=True)
    # имя в архиве распаковано в escaped-виде (#U0410...), поэтому кладём файл
    # имя файла роли не играет: stage берёт единственный xlsx папки склада
    assort = glob.glob(str(PROJECT / "input" / "mvideo" / "*.xlsx"))[0]
    shutil.copy(assort, ctx.resolved_input_dir() / pathlib.Path(assort).name)

    ctx.psb_session = lambda: None
    ctx.admin_session = lambda: None
    psb_mod.fetch_admin_category = lambda s, b, cid: FAKE_CATEGORY

    def fake_find(session, base, offer_id, seller_id, **kw):
        return {"id": f"int-{offer_id}", "offerId": offer_id}

    def fake_detail(session, base, product_id, seller_id):
        offer_id = str(product_id).replace("int-", "")
        # чередуем категории; у части товаров нет штрихкода и описания —
        # именно они должны попасть в отчёт и получить заглушки
        idx = int(offer_id) % 3
        detail = {
            "offerId": offer_id, "name": f"Товар {offer_id}",
            "price": 1200, "oldPrice": 1500, "vat": 7, "quantity": 3,
            "weight": 2.0, "width": 10, "height": 20, "length": 30,
            "primaryImage": "https://example.test/1.jpg",
            "images": ["https://example.test/2.jpg", "https://example.test/3.jpg"],
            "categoryId": psb_ids[idx],
            "attributes": [{"attributeName": "Бренд", "attributeValue": "TestBrand"}],
        }
        if idx != 0:
            detail["barcode"] = "4601234567890"
            detail["description"] = "описание от apigw"
        return detail

    psb_mod.find_product_by_offer_id = fake_find
    psb_mod.get_product_details = fake_detail
    return ctx, psb_ids


def _run_psb(tmp: Path):
    ctx, psb_ids = _fake_psb_env(tmp)
    buf = io.StringIO()
    with redirect_stdout(buf):
        result = psb_mod.PsbGenerateStage().run(ctx)
    return ctx, psb_ids, result, buf.getvalue()


def test_mvideo_template_matches_bazar_shape():
    tmp = Path(tempfile.mkdtemp()); cwd = os.getcwd()
    try:
        os.chdir(tmp)
        ctx, psb_ids, result, out = _run_psb(tmp)
        assert result.ok, result

        files = sorted(ctx.resolved_output_dir().glob("catalog_*.xlsx"))
        assert files, "шаблоны не созданы"

        bridge = json.load(open("console_to_admin_category_map.json", encoding="utf-8"))
        for f in files:
            wb = load_workbook(f, data_only=True)
            hd = [c.value for c in wb["Products"][1]]

            # 1) B2B-колонки есть во всех файлах
            for col in B2B:
                assert col in hd, f"{f.name}: нет колонки {col}"

            # 2) LockedInfo — id ПСБ (ключ моста), а не admin_id (значение моста)
            a1 = wb["LockedInfo"]["A1"].value
            assert str(a1) in bridge, f"{f.name}: LockedInfo={a1} не ключ моста (похоже на admin_id)"
            assert a1 not in set(bridge.values()) or str(a1) in bridge
            assert str(a1) in f.name, f"{f.name}: имя файла и LockedInfo разошлись"

            ws = wb["Products"]
            i_extra = hd.index("Дополнительное Изображения (до 10 штук)")
            i_bc = hd.index("Штрихкод *")
            i_tnved = hd.index("ТН ВЭД *")
            for r in ws.iter_rows(min_row=2, values_only=True):
                if all(v in (None, "") for v in r):
                    continue
                # 3) разделитель изображений — как у Базара
                extra = str(r[i_extra] or "")
                if extra:
                    assert ";" not in extra and IMAGES_SEPARATOR in extra, extra
                # 4) штрихкод ровно 13 символов
                assert len(str(r[i_bc])) == 13, (f.name, r[i_bc])
                # 5) B2B-заглушка проставлена
                assert str(r[i_tnved]) == "8500000000", r[i_tnved]
        print(f"    файлов {len(files)}, LockedInfo из пространства ПСБ: "
              f"{[load_workbook(f)['LockedInfo']['A1'].value for f in files]}")
    finally:
        os.chdir(cwd); shutil.rmtree(tmp, ignore_errors=True)


def test_report_lists_products_and_asks():
    tmp = Path(tempfile.mkdtemp()); cwd = os.getcwd()
    try:
        os.chdir(tmp)
        ctx, _, result, out = _run_psb(tmp)
        assert "ОБЯЗАТЕЛЬНЫЕ ПОЛЯ БЕЗ ДАННЫХ" in out, "нет общего отчёта о пропусках"
        assert "Настроенные заглушки" in out, "не показан список заглушек"
        assert "Подставить заглушки" in out, "не задан вопрос"
        assert "Id " in out, "в отчёте нет списка товаров"
        report = Path(result.summary["missing_report"])
        assert report.is_file(), report
        body = report.read_text(encoding="utf-8-sig")
        assert "Id Товара" in body and body.count("\n") > 1
        applied = result.summary["placeholders_applied"]
        assert applied.get("ТН ВЭД"), applied
        assert result.summary["placeholders_declined"] is False
        print(f"    подставлено: {applied}")
    finally:
        os.chdir(cwd); shutil.rmtree(tmp, ignore_errors=True)


def test_decline_leaves_cells_empty():
    tmp = Path(tempfile.mkdtemp()); cwd = os.getcwd()
    try:
        os.chdir(tmp)
        ctx, _, _ = _fake_psb_env(tmp)[0], None, None
        ctx.assume_yes = False          # неинтерактивный прогон без --yes = «нет»
        buf = io.StringIO()
        with redirect_stdout(buf):
            result = psb_mod.PsbGenerateStage().run(ctx)
        assert result.summary["placeholders_declined"] is True
        assert not result.summary["placeholders_applied"]

        empty_found = False
        for f in ctx.resolved_output_dir().glob("catalog_*.xlsx"):
            wb = load_workbook(f, data_only=True); ws = wb["Products"]
            hd = [c.value for c in ws[1]]; i = hd.index("ТН ВЭД *")
            for r in ws.iter_rows(min_row=2, values_only=True):
                if any(v not in (None, "") for v in r):
                    assert r[i] in (None, ""), "заглушка подставлена вопреки отказу"
                    empty_found = True
        assert empty_found
    finally:
        os.chdir(cwd); shutil.rmtree(tmp, ignore_errors=True)


def test_attribute_mapping_file_per_warehouse():
    """Файл маппинга свой у каждого склада; Базар остаётся на историческом."""
    ctxs = {m.CONFIG.key: PipelineContext(config=m.CONFIG, env={}, interactive=False, assume_yes=True)
            for m in (bazar, mvideo)}
    assert ctxs["bazar"].resolved_attribute_mapping_file() == Path("attribute_mapping.json")
    assert ctxs["mvideo"].resolved_attribute_mapping_file() == Path("attribute_mapping_mvideo.json")
    assert ctxs["bazar"].resolved_attribute_mapping_file() != ctxs["mvideo"].resolved_attribute_mapping_file()
    # исторический файл Базара читается и не пуст
    if Path(PROJECT / "attribute_mapping.json").is_file():
        data = json.load(open(PROJECT / "attribute_mapping.json", encoding="utf-8"))
        assert len(data) > 50, len(data)
        brand = data.get("Бренд(Brend)", {})
        names = brand.get("source_names", brand.get("bazar_names", []))
        assert "Бренд одежды" in names and "Бренды косметики" in names, names


def test_no_empty_stub_written_when_not_prompting():
    """Неинтерактивный прогон больше не пишет в файл пустую заглушку —
    иначе атрибут навсегда считался бы «замапленным на пустоту»."""
    from belka.core.attr_mapping import AttributeMapper
    tmp = Path(tempfile.mkdtemp())
    try:
        path = tmp / "attribute_mapping_test.json"
        mapper = AttributeMapper(path, interactive=False, prompt_for_missing=True)
        assert mapper.ensure_required("Материал", {"Материал одежды"}, ["ctx"]) is None
        assert not path.exists(), "файл не должен создаваться пустой заглушкой"
        assert "Материал" in mapper.report()["unresolved_required"]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_mapping_is_fallback_after_direct_match():
    """Порядок: прямое -> нормализованное -> имена из файла -> дефолт."""
    from belka.core.attr_mapping import AttributeMapper
    tmp = Path(tempfile.mkdtemp())
    try:
        path = tmp / "m.json"
        path.write_text(json.dumps({
            "Бренд(Brend)": {"source_names": ["Бренд одежды"], "default": "Без бренда"}
        }, ensure_ascii=False), encoding="utf-8")
        mapper = AttributeMapper(path, interactive=False, prompt_for_missing=False)

        def lookup_for(values):
            def lookup(name, normalized=False):
                if normalized:
                    return {normalize_attr_key(k): v for k, v in values.items()}.get(name)
                return values.get(name)
            return lookup

        # 1) прямое совпадение имени выигрывает у маппинга
        assert mapper.value_for("Бренд(Brend)", lookup_for({"Бренд(Brend)": "Acer"})) == "Acer"
        # 2) нормализованное — тоже раньше маппинга
        assert mapper.value_for("Бренд(Brend)", lookup_for({"Бренд": "Huawei"})) == "Huawei"
        # 3) имя из файла маппинга
        assert mapper.value_for("Бренд(Brend)", lookup_for({"Бренд одежды": "Avesta"})) == "Avesta"
        # 4) дефолт — и он посчитан
        assert mapper.value_for("Бренд(Brend)", lookup_for({})) == "Без бренда"
        assert mapper.report()["defaults_used"] == 1
        # нет записи в файле — пусто, а не выдумка
        assert mapper.value_for("Цвет", lookup_for({})) == ""
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    print("=== общее ===")
    check("штрихкод ровно 13 символов", test_barcode_exactly_13)
    check("нормализация имени атрибута", test_normalize_attr_key)
    check("реестр заглушек и переопределение складом", test_placeholder_registry_and_override)
    check("колонки реестра совпадают с колонками шаблона", test_every_placeholder_column_is_known)

    print("\n=== Мвидео на реальном файле ассортимента ===")
    check("шаблон совпал по форме с Базаром", test_mvideo_template_matches_bazar_shape)
    check("отчёт со списком товаров и вопрос", test_report_lists_products_and_asks)
    check("отказ оставляет ячейки пустыми", test_decline_leaves_cells_empty)

    print("\n=== маппинг атрибутов по складам ===")
    check("файл маппинга свой у каждого склада", test_attribute_mapping_file_per_warehouse)
    check("пустая заглушка в файл не пишется", test_no_empty_stub_written_when_not_prompting)
    check("маппинг — запасной путь после совпадения имени", test_mapping_is_fallback_after_direct_match)

    failed = [n for n, ok in RESULTS if not ok]
    print(f"\n=== ИТОГ: {len(RESULTS) - len(failed)}/{len(RESULTS)} ===")
    sys.exit(1 if failed else 0)
