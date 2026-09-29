#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты на ротацию изображений (core/images.py::rotate_images) и снятие HTML
(core/text.py::strip_html), плюс прогон psb_generate Мвидео целиком на
заглушках apigw/admin — без сети.

Запуск:  python3 tests/test_images_and_html.py
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import traceback
from pathlib import Path

PROJECT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT))

import openpyxl                                                          # noqa: E402

from belka.context import PipelineContext                                # noqa: E402
from belka.core import images as img                                     # noqa: E402
from belka.core.images import rotate_images                              # noqa: E402
from belka.core.text import strip_html                                   # noqa: E402
from belka.stages import psb_generate as psb_mod                         # noqa: E402
from belka.warehouses import mvideo, b2c, bazar                          # noqa: E402

RESULTS = []


def check(name, fn):
    try:
        fn()
    except Exception:
        RESULTS.append((name, False))
        print(f"✘ {name}")
        traceback.print_exc()
    else:
        RESULTS.append((name, True))
        print(f"✔ {name}")


# --- rotate_images -------------------------------------------------------------

def test_single_main_unchanged():
    r = rotate_images("m", [])
    assert (r.action, r.main_after, r.additional_after, r.removed) == (img.SINGLE_MAIN_UNCHANGED, "m", [], [])
    assert not r.changed


def test_single_additional_removed_main_kept():
    r = rotate_images("m", ["a1"])
    assert r.action == img.SINGLE_ADDITIONAL_REMOVED
    assert r.main_after == "m" and r.additional_after == []
    assert r.removed == [("доп.", 1, "a1")]


def test_two_additional():
    r = rotate_images("m", ["a1", "a2"])
    assert r.action == img.ROTATED
    assert r.main_after == "a1" and r.additional_after == []
    assert r.removed == [("основное", 1, "m"), ("доп.", 2, "a2")]
    assert r.promoted_from == 1


def test_many_additional():
    r = rotate_images("m", ["a1", "a2", "a3", "a4", "a5"])
    assert r.main_after == "a1" and r.additional_after == ["a2", "a3", "a4"]
    assert r.removed[-1] == ("доп.", 5, "a5")
    line = r.describe()
    assert "было осн. 1 + доп. 5" in line and "доп. #5 из 5" in line and "стало осн. 1 + доп. 3" in line, line


def test_duplicate_of_main_removed_before_rotation():
    # если apigw кладёт основное ещё и первым в images — без чистки дубля
    # «новым основным» стало бы то же самое удаляемое фото
    r = rotate_images("m", ["m", "a1", "", "a1", "a2", "a3"])
    assert r.main_after == "a1" and r.additional_after == ["a2"]
    assert [p for p, _ in r.duplicates] == [1, 3, 4]
    assert r.removed == [("основное", 1, "m"), ("доп.", 6, "a3")]  # № — из исходного списка
    assert r.promoted_from == 2


def test_no_main():
    r = rotate_images(None, ["a1", "a2"])
    assert r.action == img.NO_MAIN_UNCHANGED and r.main_after is None and r.additional_after == ["a1", "a2"]
    assert rotate_images("  ", None).action == img.NO_IMAGES


def test_input_not_mutated():
    add = ["a1", "a2", "a3"]
    rotate_images("m", add)
    assert add == ["a1", "a2", "a3"]


def test_extract_images_bazar_unchanged():
    assert img.extract_images(["a", "b", "a", "c"]) == ("a", "b, c")
    assert img.extract_images([]) == (None, None)


# --- strip_html ------------------------------------------------------------------

def test_plain_text_untouched():
    for v in ("Обычный текст", "Размер < 5 см и > 2 см", "", None, 42):
        assert strip_html(v) == v, v


def test_basic_tags_and_breaks():
    s = "<p>Первый <b>абзац</b>.</p><p>Второй<br>строка<br/>ещё</p>"
    assert strip_html(s) == "Первый абзац.\n\nВторой\nстрока\nещё", repr(strip_html(s))


def test_lists():
    s = "Плюсы:<ul><li>тихий</li><li>лёгкий</li></ul>Итог"
    assert strip_html(s) == "Плюсы:\n• тихий\n• лёгкий\nИтог", repr(strip_html(s))


def test_entities_and_nbsp():
    assert strip_html("Цена&nbsp;100&nbsp;&#8381; &amp; &laquo;скидка&raquo;") == "Цена 100 ₽ & «скидка»"


def test_script_style_removed():
    assert strip_html("<style>p{color:red}</style>Текст<script>alert(1)</script>") == "Текст"


def test_escaped_markup():
    assert strip_html("&lt;p&gt;Текст&lt;/p&gt;") == "Текст"


def test_only_markup_gives_none():
    assert strip_html("<p> &nbsp; </p><br>") is None


def test_attributes_and_nesting():
    s = '<div class="x"><span style="a:b">Экран <i>6,1"</i></span></div>'
    assert strip_html(s) == 'Экран 6,1"'


# --- флаги складов ----------------------------------------------------------------

def test_flags_only_mvideo():
    assert mvideo.CONFIG.rotate_images and mvideo.CONFIG.strip_html_description
    for cfg in (b2c.CONFIG, bazar.CONFIG):
        assert not cfg.rotate_images and not cfg.strip_html_description, cfg.key


# --- psb_generate целиком, на заглушках ----------------------------------------------

FAKE_CATEGORY = {"properties": [{"display_name": "Бренд", "is_required": False, "has_directory": True}]}

FAKE_DETAILS = {
    "A1": {"description": "<p>Мощный <b>пылесос</b></p><ul><li>тихий</li></ul>",
           "primaryImage": "m1", "images": ["a1", "a2", "a3", "a4"]},
    "A2": {"description": "Без разметки", "primaryImage": "m2", "images": []},
    "A3": {"description": "<br>&nbsp;", "primaryImage": "m3", "images": ["b1"]},
}


def _fake_detail(offer_id):
    base = {"offerId": offer_id, "name": f"Товар {offer_id}", "barcode": f"460000000000{offer_id[-1]}",
            "categoryId": 26014, "price": 1200, "vat": 4, "quantity": 5,
            "weight": 1, "width": 1, "height": 1, "length": 1, "attributes": []}
    return {**base, **FAKE_DETAILS[offer_id]}


def test_psb_generate_mvideo_end_to_end():
    workdir = Path(tempfile.mkdtemp())
    cwd = os.getcwd()
    saved = (psb_mod.find_product_by_offer_id, psb_mod.get_product_details, psb_mod.fetch_admin_category)
    try:
        shutil.copy(PROJECT / "console_to_admin_category_map.json", workdir)
        os.chdir(workdir)
        ctx = PipelineContext(config=mvideo.CONFIG, env={}, interactive=False, assume_yes=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Материал", "Розничная цена", "Остатки"])
        for oid in FAKE_DETAILS:
            ws.append([oid, None, None])
        wb.save(ctx.resolved_input_dir() / "assortment.xlsx")

        ctx.psb_session = lambda: None
        ctx.admin_session = lambda: None
        psb_mod.find_product_by_offer_id = lambda s, base, oid, seller: {"id": oid}
        psb_mod.get_product_details = lambda s, base, pid, seller: _fake_detail(pid)
        psb_mod.fetch_admin_category = lambda s, base, cid: FAKE_CATEGORY

        result = psb_mod.PsbGenerateStage().run(ctx)
        s = result.summary
        assert result.ok
        assert (s["images_rotated"], s["images_single_main_unchanged"], s["images_single_additional_removed"]) == (1, 1, 1), s
        assert s["images_removed_total"] == 3, s          # A1: осн.+посл.доп., A3: одно доп.
        assert s["html_descriptions_cleaned"] == 2, s     # A1 и A3 (у A2 разметки нет)

        out = Path(s["output_dir"])
        cat = openpyxl.load_workbook(next(out.glob("catalog_*.xlsx")))["Products"]
        headers = [c.value for c in cat[1]]
        rows = {r[headers.index("Id Товара *")]: r for r in cat.iter_rows(min_row=2, values_only=True)}
        col = lambda name: headers.index(name)                                     # noqa: E731
        main, add, desc = (col("Основное изображение *"),
                           col("Дополнительное Изображения (до 10 штук)"), col("Описание *"))

        assert rows["A1"][main] == "a1" and rows["A1"][add] == "a2, a3", rows["A1"]
        assert rows["A1"][desc] == "Мощный пылесос\n\n• тихий", repr(rows["A1"][desc])
        assert rows["A2"][main] == "m2" and rows["A2"][add] is None
        assert rows["A2"][desc] == "Без разметки"
        assert rows["A3"][main] == "m3" and rows["A3"][add] is None
        assert rows["A3"][desc] == "нет описания", rows["A3"][desc]   # пусто после HTML -> заглушка

        # отчёт не должен лежать рядом с каталогами — import/upload берут все *.xlsx папки
        assert [p.name for p in out.glob("*.xlsx")] == [p.name for p in out.glob("catalog_*.xlsx")]
        rep = openpyxl.load_workbook(out / "reports" / "images_report_mvideo.xlsx").active
        rep_rows = list(rep.iter_rows(min_row=2, values_only=True))
        assert [r[0] for r in rep_rows] == ["A1", "A2", "A3"]
        assert rep_rows[0][4] == "основное #1, доп. #4" and rep_rows[0][5] == 1, rep_rows[0]
        assert rep_rows[1][2:4] == (1, 0) and rep_rows[1][4] is None
        assert rep_rows[2][4] == "доп. #1"
    finally:
        psb_mod.find_product_by_offer_id, psb_mod.get_product_details, psb_mod.fetch_admin_category = saved
        os.chdir(cwd)
        shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    print("=== rotate_images ===")
    check("одна картинка в основном — без изменений", test_single_main_unchanged)
    check("одно доп. — удаляется, основное остаётся", test_single_additional_removed_main_kept)
    check("два доп. — основное <- доп.1, доп. пусто", test_two_additional)
    check("пять доп. — основное <- доп.1, доп. = 2..4", test_many_additional)
    check("дубль основного в доп. убирается до ротации", test_duplicate_of_main_removed_before_rotation)
    check("нет основного — без изменений", test_no_main)
    check("входной список не мутируется", test_input_not_mutated)
    check("extract_images Базара не изменилась", test_extract_images_bazar_unchanged)

    print("\n=== strip_html ===")
    check("текст без разметки не трогается", test_plain_text_untouched)
    check("теги и переносы", test_basic_tags_and_breaks)
    check("списки", test_lists)
    check("сущности и nbsp", test_entities_and_nbsp)
    check("script/style удаляются с содержимым", test_script_style_removed)
    check("экранированная разметка", test_escaped_markup)
    check("одна разметка -> None", test_only_markup_gives_none)
    check("атрибуты и вложенность", test_attributes_and_nesting)

    print("\n=== склады ===")
    check("флаги включены только у Мвидео", test_flags_only_mvideo)
    check("psb_generate Мвидео целиком (заглушки)", test_psb_generate_mvideo_end_to_end)

    failed = [n for n, ok in RESULTS if not ok]
    print(f"\n=== ИТОГ: {len(RESULTS) - len(failed)}/{len(RESULTS)} ===")
    sys.exit(1 if failed else 0)
