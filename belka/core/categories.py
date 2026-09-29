#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Загрузка и разбор дерева категорий — общее для админки и консоли (оба
хоста отдают один и тот же формат через POST .../categories:tree).

fetch_admin_category / build_template_columns / write_catalog_file — раньше
были приватными методами GenerateStage (Bazar). Вынесены сюда (09.2026),
когда понадобился второй источник товаров с ТЕМ ЖЕ форматом шаблона Белки
(belka/stages/psb_generate.py — склады №6/№7): оба генератора строят колонки
xlsx по ОДНОЙ и той же категорийной схеме Белки (?include=properties,
hidden_properties), различаются только тем, откуда берут сами товары и как
заполняют ячейки."""
from __future__ import annotations

from pathlib import Path
from datetime import datetime

import openpyxl

from .xlsx import build_header

CATEGORY_URL_PART = "/api/v1/catalog/categories/{id}"
CATEGORY_INCLUDE = "properties,hidden_properties"


def fetch_admin_category(session, api_base: str, category_id) -> dict:
    """GET .../categories/{id}?include=properties,hidden_properties — карточка
    категории Белки со списком атрибутов (properties), т.е. что реально
    требуется на этот момент, а не что было зашито когда-то в статический
    шаблон."""
    url = f"{api_base}{CATEGORY_URL_PART.format(id=category_id)}"
    resp = session.get(url, params={"include": CATEGORY_INCLUDE}, timeout=30)
    return resp.json()["data"]


def build_template_columns(category_data: dict, direct_fields: list) -> list:
    """direct_fields — список (header, is_required, filler) фиксированных
    структурных колонок (не зависят от категории, у каждого генератора свои
    filler'ы под форму своего источника товаров). Дополняется атрибутами
    категории (category_data['properties']) через build_header(). Если
    property называется так же, как уже занятый заголовок из direct_fields —
    пропускается (структурная колонка в приоритете)."""
    columns = []
    for header, is_required, filler in direct_fields:
        columns.append({"header": header, "is_required": is_required, "has_directory": False,
                         "kind": "direct", "filler": filler})
    for prop in category_data.get("properties", []):
        display_name = prop.get("display_name") or prop.get("name")
        is_required = bool(prop.get("is_required"))
        has_directory = bool(prop.get("has_directory"))
        header = build_header(display_name, is_required, has_directory)
        if any(c["header"] == header for c in columns):
            continue
        columns.append({"header": header, "is_required": is_required, "has_directory": has_directory,
                         "kind": "attribute", "display_name": display_name})
    return columns


def write_catalog_file(columns: list, rows: list, category_id, out_path: Path) -> None:
    """Пишет один xlsx-каталог: лист 'Products' с колонками+строками и
    служебный лист 'LockedInfo' с id категории Белки в A1 (это то, что потом
    читает stage 'bridge' Bazar при сверке путей — не трогать формат)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    for col_idx, col in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col["header"])
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    locked = wb.create_sheet("LockedInfo")
    locked["A1"] = category_id
    wb.save(out_path)


def default_catalog_filename(category_id) -> str:
    """catalog_<id>_<yy-mm-dd-hh>.xlsx — общее имя выходного файла что у
    Bazar, что у ПСБ-складов (по одному файлу на категорию Белки за прогон)."""
    date_str = datetime.now().strftime("%y-%m-%d-%H")
    return f"catalog_{category_id}_{date_str}.xlsx"


def clean_stale_catalog_files(output_dir: Path) -> int:
    """Удаляет все catalog_*.xlsx в папке вывода ПЕРЕД генерацией новых —
    вызывается в начале run() у GenerateStage/PsbGenerateStage, до первой
    записи файла.

    09.2026 (по просьбе пользователя, реальный случай): resolved_output_dir()
    — папка НА ДЕНЬ, а не на прогон (output_dir/<склад>/ГГГГ-ММ-ДД), а имя
    файла (default_catalog_filename) включает лишь ЧАС, а не полное время —
    поэтому второй прогон в тот же день (например, после правки бага) не
    заменяет файлы первого прогона, а кладёт новые рядом, если час
    отличается. В папке накапливаются вперемешку старые (до правки) и новые
    (после правки) catalog_*.xlsx одного и того же дня — import_check/upload
    сканируют папку целиком и видят проблемы из СТАРЫХ файлов тоже, из-за
    чего кажется, что фикс не подействовал, хотя новые файлы уже верны.

    Трогает ТОЛЬКО catalog_*.xlsx. Остальные файлы в output_dir
    (upload_log.json, added_values.json, deactivated_by_certificate.json,
    category_bridge_unmatched.csv и т.п.) — это персистентное состояние
    (история загрузок, ранее принятые решения), а не сгенерированные
    каталоги, и НЕ удаляются."""
    removed = 0
    for path in output_dir.glob("catalog_*.xlsx"):
        path.unlink()
        removed += 1
    return removed


def fetch_category_tree(session, api_base: str, url_part: str = "/api/v1/catalog/categories:tree") -> dict:
    """POST .../categories:tree — возвращает сырой ответ {'data': [...]}."""
    url = f"{api_base}{url_part}"
    body = {"sort": [], "filter": {}}
    resp = session.post(url, json=body, timeout=30)
    return resp.json()


def extract_leaf_categories(tree_response: dict) -> list:
    """Разворачивает дерево категорий в список листовых: [{id, path, is_hidden,
    is_active, is_real_active}, ...], path — полный путь через ' / '."""
    def walk(category: dict, path: str = "") -> list:
        name = (category.get("name") or "").strip()
        new_path = f"{path} / {name}" if path else name
        children = category.get("children") or []
        if not children:
            return [{
                "id": category.get("id"),
                "path": new_path,
                "is_hidden": category.get("is_hidden"),
                "is_active": category.get("is_active"),
                "is_real_active": category.get("is_real_active"),
            }]
        leaves = []
        for child in children:
            leaves.extend(walk(child, new_path))
        return leaves

    all_leaves = []
    for category in tree_response.get("data") or []:
        all_leaves.extend(walk(category))
    return all_leaves