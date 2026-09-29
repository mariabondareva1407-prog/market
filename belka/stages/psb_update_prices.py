#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Stage 'psb_update_prices' — ВРЕМЕННЫЙ (по решению пользователя, 09.2026)
способ подготовки шаблонов для склада №7 (b2c): в отличие от Мвидео, шаблоны
для b2c не генерируются с нуля через категорийную API Белки (см.
psb_generate.py) — они УЖЕ готовы и лежат в input_dir (обычно просто
'input', см. WarehouseConfig.input_dir), заполнены целиком (название,
атрибуты, категория и т.д.), и требуют только обновления трёх ячеек: "Цена,
руб *", "Цена до скидки, руб", "Количество Товара *". Остальное не трогаем.

Прямой перенос присланного update_prices.py на движок Stage, с той же
поправкой, что и у psb_generate.py — sellerId в параметрах поиска
(find_product_by_offer_id уже сделан правильно в core/psb_api.py). Разница
с исходным скриптом: там цена/цена до скидки копируются из apigw как есть;
здесь, по решению пользователя, к ним ДОПОЛНИТЕЛЬНО применяется та же
арифметика вычитания НДС, что и в psb_generate.py (см. core/psb_api.py:
strip_vat) — чтобы сайт при импорте не задвоил НДС собственным умножением
на (1 + ндс%/100) * 1.01.

Читает шаблоны из cfg.input_dir, пишет ОБНОВЛЁННЫЕ КОПИИ в cfg.output_dir
(исходные файлы в input_dir не трогает) — output_dir и есть дальше вход для
import_check/import_attributes/upload, как и у остальных складов.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..core.psb_api import find_product_by_offer_id, strip_vat
from ..core.http import SessionBlockedError
from ..pipeline import Stage, StageResult

PRODUCTS_SHEET_NAME = "Products"

COL_ID_TOVARA = "Id Товара *"
COL_PRICE = "Цена, руб *"
COL_OLD_PRICE = "Цена до скидки, руб"
COL_QUANTITY = "Количество Товара *"

REQUIRED_HEADERS = {COL_ID_TOVARA, COL_PRICE, COL_OLD_PRICE, COL_QUANTITY}


def _find_header_columns(ws) -> dict:
    headers = {}
    for cell in ws[1]:
        if cell.value:
            headers[cell.value] = cell.column
    return headers


class PsbUpdatePricesStage(Stage):
    name = "psb_update_prices"

    def __init__(self, folder=None):
        self.folder_override = Path(folder) if folder else None

    def _process_workbook(self, session, base_url, seller_id, in_path: Path, out_path: Path) -> dict:
        wb = load_workbook(in_path)
        if PRODUCTS_SHEET_NAME not in wb.sheetnames:
            print(f"  [!] {in_path.name} пропущен: нет листа '{PRODUCTS_SHEET_NAME}'")
            return {"updated": 0, "not_found": 0, "skipped_file": True}

        ws = wb[PRODUCTS_SHEET_NAME]
        headers = _find_header_columns(ws)
        missing = REQUIRED_HEADERS - headers.keys()
        if missing:
            print(f"  [!] {in_path.name} пропущен: не найдены колонки {sorted(missing)}")
            return {"updated": 0, "not_found": 0, "skipped_file": True}

        col_id, col_price = headers[COL_ID_TOVARA], headers[COL_PRICE]
        col_old_price, col_qty = headers[COL_OLD_PRICE], headers[COL_QUANTITY]

        updated = 0
        not_found = 0
        row_num = 2
        while True:
            offer_id_cell = ws.cell(row=row_num, column=col_id).value
            if offer_id_cell in (None, ""):
                break  # конец данных
            offer_id = str(int(offer_id_cell)) if isinstance(offer_id_cell, float) else str(offer_id_cell)

            product = find_product_by_offer_id(session, base_url, offer_id, seller_id)
            if not product:
                print(f"  [!] {in_path.name}, строка {row_num}: offer_id={offer_id} не найден в API, строка не тронута")
                not_found += 1
                row_num += 1
                continue

            vat_code = product.get("vat")
            ws.cell(row=row_num, column=col_price, value=strip_vat(product.get("price"), vat_code))
            ws.cell(row=row_num, column=col_old_price, value=strip_vat(product.get("oldPrice"), vat_code))
            ws.cell(row=row_num, column=col_qty, value=product.get("quantity"))
            updated += 1
            row_num += 1

        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
        return {"updated": updated, "not_found": not_found, "skipped_file": False}

    def run(self, ctx) -> StageResult:
        cfg = ctx.config
        if not cfg.psb_api_base:
            raise RuntimeError(f"У склада '{cfg.key}' не задан psb_api_base.")
        psb = ctx.psb_session()

        seller_id_psb = int(ctx.override("PSB_SELLER_ID") or cfg.psb_seller_id or 0)
        if not seller_id_psb:
            raise RuntimeError(f"У склада '{cfg.key}' не задан psb_seller_id (id продавца на стороне ПСБ).")

        # Единственный stage, который берёт ВСЕ xlsx папки, а не один:
        # у b2c вход — пачка готовых шаблонов, каждый обновляется отдельно.
        in_folder = self.folder_override or ctx.resolved_input_dir()  # input_dir/b2c
        files = sorted(in_folder.glob("*.xlsx")) if self.folder_override else ctx.input_files(".xlsx")
        if not files:
            raise RuntimeError(
                f"В папке {in_folder} нет .xlsx шаблонов для обновления — положите готовые "
                f"шаблоны b2c в эту папку."
            )

        out_folder = ctx.resolved_output_dir()  # output_dir/b2c/ГГГГ-ММ-ДД, создаётся сама
        print(f"[psb_update_prices] шаблонов к обновлению: {len(files)} (из {in_folder}, "
              f"обновлённые копии -> {out_folder})")

        total_updated = 0
        total_not_found = 0
        skipped_files = 0
        blocked = False
        for i, path in enumerate(files, 1):
            print(f"  [{i}/{len(files)}] {path.name}")
            try:
                result = self._process_workbook(psb, cfg.psb_api_base, seller_id_psb, path, out_folder / path.name)
            except SessionBlockedError as e:
                print(f"\n⛔ ОСТАНОВКА: {e}")
                print("Обновление прервано — обновите токен и перезапустите (уже сохранённые файлы не потеряны).")
                blocked = True
                break
            if result["skipped_file"]:
                skipped_files += 1
                continue
            total_updated += result["updated"]
            total_not_found += result["not_found"]
            print(f"    обновлено строк: {result['updated']}, не найдено в API: {result['not_found']}")

        summary = {
            "files": len(files),
            "skipped_files": skipped_files,
            "rows_updated": total_updated,
            "rows_not_found": total_not_found,
            "output_dir": str(out_folder),
            "blocked": blocked,
        }
        print("\n=== ИТОГ psb_update_prices ===")
        for k, v in summary.items():
            print(f"  {k}: {v}")
        return StageResult(self.name, ok=not blocked, blocked=blocked, summary=summary)
