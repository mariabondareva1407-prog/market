#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Общее для stage'ов, на которые разобран прежний единый 'import':
import_check.py / import_attributes.py (был ещё import_capitalize.py, убран
09.2026 — см. belka/warehouses/bazar.py). Это НЕ stage сам по себе — только
то, что им одинаково нужно: сканирование xlsx-шаблона в общую структуру
данных и мелкие текстовые хелперы.

Каждый stage сканирует файлы заново, а не получает готовые данные от
соседа — тот же принцип, что и везде в пайплайне (источник истины между
stage'ами — файлы на диске, а не общая память в процессе). Это и даёт им
независимость: любой можно прогнать через --only отдельно от других.
"""
from __future__ import annotations

import re
import difflib
from pathlib import Path
from collections import Counter, defaultdict

import openpyxl

from ..core.text import casefold_ru
from ..core.xlsx import find_products_sheet, clean_header

SHEET_NAME_CANDIDATES = ("Products", "products", "Товары")
ID_COLUMN_HINTS = ("id товара",)
QUANTITY_COLUMN_HINTS = ("количество товара", "количество")

RARE_MIN_MAX_COUNT = 8           # порог "доминирующего" значения в колонке (эвристика "редкое")
SIMILARITY_THRESHOLD = 0.86      # порог схожести для поиска опечаток/дублей
ATTR_NAME_MATCH_THRESHOLD = 0.8  # порог схожести значения с названием атрибута

PLACEHOLDERS = {
    "-", "—", "н/д", "нет", "нет данных", "не указано",
    "test", "тест", "n/a", "null", "0",
}
STOPWORDS_FOR_ATTR_MATCH = {"женский", "мужской", "детский", "обуви", "размер", "вид", "материал", "тип"}

REPORT_SUBDIR = "report"


def normalize_capitalize(value: str) -> str:
    v = re.sub(r"\s+", " ", str(value).strip())
    if not v:
        return v
    return v[0].upper() + v[1:]


def is_quantity_header(header) -> bool:
    if not header:
        return False
    h = str(header).lower()
    return any(hint in h for hint in QUANTITY_COLUMN_HINTS)


def split_multi(raw_value) -> list:
    if raw_value is None:
        return []
    parts = [p.strip() for p in str(raw_value).split(";")]
    return [p for p in parts if p]


def is_placeholder(value_norm: str) -> bool:
    return casefold_ru(value_norm) in PLACEHOLDERS


def looks_like_attr_name(value_norm: str, attr_name: str) -> bool:
    val_l = casefold_ru(value_norm)
    if len(val_l.split()) != 1:
        return False
    words = [casefold_ru(w) for w in attr_name.split() if len(w) >= 4]
    for w in words:
        if w in STOPWORDS_FOR_ATTR_MATCH:
            continue
        if difflib.SequenceMatcher(None, val_l, w).ratio() >= ATTR_NAME_MATCH_THRESHOLD:
            return True
    return False


def find_near_duplicates(value_norm: str, candidates: list):
    target = casefold_ru(value_norm)
    for c in candidates:
        c_cf = casefold_ru(c)
        if c_cf == target:
            continue
        if difflib.SequenceMatcher(None, target, c_cf).ratio() >= SIMILARITY_THRESHOLD:
            return c
    return None


def log_key(property_id, attr_name: str) -> str:
    return f"{property_id}, {attr_name}"


def merge_into_log(log: dict, key: str, values: list) -> None:
    bucket = log.setdefault(key, [])
    existing_cf = {casefold_ru(v) for v in bucket}
    for v in values:
        if casefold_ru(v) not in existing_cf:
            bucket.append(v)
            existing_cf.add(casefold_ru(v))


def scan_file(path: Path) -> dict:
    """Читает один xlsx-шаблон ровно один раз и достаёт из него всё, что
    может понадобиться import_*-stage'ам: артикулы (для сверки
    пропавших/дублей в import_check), значения атрибутов-справочников (для
    import_attributes), количество строк с нулевым остатком (метрика в
    import_check — сам остаток теперь никто не правит, см. bazar.py)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = find_products_sheet(wb, sheet_name_candidates=SHEET_NAME_CANDIDATES, header_hints=ID_COLUMN_HINTS)

    headers_info = []
    id_col = None
    for c in range(1, ws.max_column + 1):
        raw_header = ws.cell(row=1, column=c).value
        name, is_required, is_directory = clean_header(raw_header)
        is_id = bool(raw_header) and any(h in str(raw_header).lower() for h in ID_COLUMN_HINTS)
        is_qty = is_quantity_header(raw_header)
        headers_info.append((c, raw_header, name, is_required, is_directory, is_id, is_qty))
        if is_id:
            id_col = c

    last_row = 1
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c, *_ in headers_info):
            last_row = r

    value_counter = defaultdict(Counter)
    missing_required = Counter()
    articles = []
    product_count = 0
    quantity_fixed_rows = 0

    for r in range(2, last_row + 1):
        row_has_data = any(ws.cell(row=r, column=c).value not in (None, "") for c, *_ in headers_info)
        if not row_has_data:
            continue
        product_count += 1
        if id_col:
            art = ws.cell(row=r, column=id_col).value
            if art not in (None, ""):
                articles.append(str(art).strip())
        for c, raw_header, name, is_required, is_directory, is_id, is_qty in headers_info:
            cell_value = ws.cell(row=r, column=c).value
            is_empty = cell_value in (None, "") or (isinstance(cell_value, str) and not cell_value.strip())
            if is_required and is_empty:
                missing_required[(str(path.name), raw_header)] += 1
            if is_directory and not is_empty:
                for token in split_multi(cell_value):
                    value_counter[name][token] += 1
            if is_qty and not is_empty:
                try:
                    qty_val = float(cell_value)
                except (TypeError, ValueError):
                    qty_val = None
                if qty_val == 0:
                    quantity_fixed_rows += 1

    return {
        "path": path, "wb": wb, "ws": ws, "headers_info": headers_info, "last_row": last_row,
        "value_counter": value_counter, "missing_required": missing_required,
        "articles": articles, "product_count": product_count,
        "quantity_fixed_rows": quantity_fixed_rows,
    }
